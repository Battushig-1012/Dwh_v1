from io import BytesIO

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from sqlalchemy import text
from sqlalchemy.orm import Session

from .. import demo_data
from ..auth.deps import require_roles
from ..config import settings
from ..db import get_db

router = APIRouter()


@router.get("/m1.xlsx")
def export_m1_excel(
    db: Session | None = Depends(get_db),
    user: dict = Depends(require_roles("REPORT", "EXEC", "AUDITOR", "RISK")),
):
    if settings.demo_mode:
        rows = demo_data.M1
    else:
        assert db is not None
        rows = db.execute(
            text(
                """
                SELECT reporting_period, dpm_code, dpm_label_mn, dpm_label_en, amount_mnt, formula_ref
                FROM data_marts.mart_reg_mongolbank_m1
                WHERE reporting_period = (SELECT MAX(reporting_period) FROM data_marts.mart_reg_mongolbank_m1)
                ORDER BY dpm_code
                """
            )
        ).mappings().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Mongolbank_M1"
    headers = ["Period", "DPM Code", "Label (MN)", "Label (EN)", "Amount MNT", "Formula"]
    header_fill = PatternFill("solid", fgColor="1B3A4B")
    header_font = Font(color="FFFFFF", bold=True)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for i, r in enumerate(rows, 2):
        ws.cell(i, 1, r["reporting_period"])
        ws.cell(i, 2, r["dpm_code"])
        ws.cell(i, 3, r["dpm_label_mn"])
        ws.cell(i, 4, r["dpm_label_en"])
        ws.cell(i, 5, float(r["amount_mnt"]))
        ws.cell(i, 6, r["formula_ref"])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 22

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    if not settings.demo_mode and db is not None:
        db.execute(
            text(
                """
                INSERT INTO governance.audit_trail (actor, action, entity_type, entity_id, after_state)
                VALUES (:actor, 'EXPORT_EXCEL', 'report', 'mongolbank_m1', '{"format":"xlsx"}'::jsonb)
                """
            ),
            {"actor": user["username"]},
        )
        db.commit()

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=mongolbank_m1.xlsx"},
    )


@router.get("/srep.pdf")
def export_srep_pdf(
    db: Session | None = Depends(get_db),
    user: dict = Depends(require_roles("REPORT", "EXEC", "AUDITOR", "RISK")),
):
    if settings.demo_mode:
        capital = demo_data.CAPITAL
        credit = demo_data.CREDIT
    else:
        assert db is not None
        capital = db.execute(
            text("SELECT * FROM data_marts.mart_srep_capital_adequacy ORDER BY as_of_date DESC LIMIT 1")
        ).mappings().first()
        credit = db.execute(
            text("SELECT * FROM data_marts.mart_srep_credit_portfolio ORDER BY as_of_date DESC LIMIT 1")
        ).mappings().first()

    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    y = height - 50
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, y, "BBSB SREP Executive Summary")
    y -= 30
    c.setFont("Helvetica", 11)
    if capital:
        c.drawString(50, y, f"As of: {capital['as_of_date']}")
        y -= 20
        c.drawString(50, y, f"CAR: {float(capital['car_pct']):.2f}% (min {float(capital['regulatory_minimum']):.1f}%)")
        y -= 20
        c.drawString(50, y, f"Total Capital: {float(capital['total_capital']):,.0f} MNT")
        y -= 20
    if credit:
        c.drawString(50, y, f"NPL Ratio: {float(credit['npl_ratio_pct']):.2f}%")
        y -= 20
        c.drawString(50, y, f"Total Outstanding: {float(credit['total_outstanding']):,.0f} MNT")
        y -= 20
    mode = "DEMO MODE" if settings.demo_mode else "LIVE"
    c.setFont("Helvetica-Oblique", 9)
    c.drawString(50, 40, f"{mode} — exported by {user['username']} — no manual edit")
    c.showPage()
    c.save()
    buf.seek(0)

    if not settings.demo_mode and db is not None:
        db.execute(
            text(
                """
                INSERT INTO governance.audit_trail (actor, action, entity_type, entity_id, after_state)
                VALUES (:actor, 'EXPORT_PDF', 'report', 'srep_summary', '{"format":"pdf"}'::jsonb)
                """
            ),
            {"actor": user["username"]},
        )
        db.commit()

    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=srep_summary.pdf"},
    )
